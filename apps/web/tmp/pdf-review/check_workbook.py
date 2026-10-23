from openpyxl import load_workbook
from pathlib import Path

book = load_workbook(Path(__file__).with_name('populated-acceptance.xlsx'))
assert book.sheetnames == ['Executive Overview', 'Detailed Report', 'SD Register']
for sheet in book:
    assert '2040' in sheet['A2'].value
    assert sheet.page_setup.orientation == 'landscape'
    print(sheet.title, 'rows', sheet.max_row, 'freeze', sheet.freeze_panes, 'fit width', sheet.page_setup.fitToWidth)
overview = book['Executive Overview']
values = {row[0].value: row[1] for row in overview.iter_rows() if row[0].value in ('Recognized sales', 'Total allocated', 'Salary Deduction')}
assert values['Recognized sales'].value == values['Total allocated'].value == 4200
assert values['Salary Deduction'].value == 450
assert all(cell.data_type == 'n' and '0.00' in cell.number_format for cell in values.values())
register = book['SD Register']
assert register['C7'].value == '000001' and register['C7'].data_type == 's'
assert register['D7'].alignment.wrap_text and register['F7'].alignment.wrap_text
assert register.auto_filter.ref == 'A6:H6' and register.freeze_panes == 'A7'
assert sum(register.cell(row, 8).value for row in range(7, 52)) == 450
print('Financial totals, numeric formats, text identifier, wrapping, filter and freeze metadata passed.')
